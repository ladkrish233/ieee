from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from models import PaperDraft

TNR = "Times New Roman"


# ---------- low-level helpers ----------

def set_two_columns(section, num_cols: int = 2, spacing: int = 360):
    """spacing is in twentieths-of-a-point; 360 = 0.25in gutter (matches template)."""
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sectPr.append(cols)
    cols.set(qn('w:num'), str(num_cols))
    cols.set(qn('w:space'), str(spacing))


def set_default_font(doc, name=TNR, size=10):
    style = doc.styles['Normal']
    style.font.name = name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def set_margins(section, top=1.0, bottom=1.0, left=0.75, right=0.75):
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)


def add_bottom_border(paragraph, size=6, color="000000"):
    """Thin horizontal rule under a paragraph — used after the keywords line."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_box_border(paragraph, color="AAAAAA", size=4):
    """Full box border — used for the figure placeholder."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        pBdr.append(el)
    pPr.append(pBdr)


def add_section_heading(doc, text: str):
    """Level-1 IEEE heading: bold, centered, ALL CAPS, 10pt."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)   # 120 twips
    p.paragraph_format.space_after = Pt(3)    # 60 twips
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = TNR
    return p


def add_subsection_a(doc, text: str):
    """Level-A subheading: bold + small caps, centered."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(9)   # 180 twips
    p.paragraph_format.space_after = Pt(3)    # 60 twips
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = TNR
    run.font.small_caps = True
    return p


def add_subsection_b(doc, text: str):
    """Level-B subheading: italic only, left-aligned."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)   # 120 twips
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.name = TNR
    return p


def add_body_paragraph(doc, text: str):
    """Standard IEEE body paragraph: first-line indent, justified, 10pt."""
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(3)  # 60 twips
    for run in p.runs:
        run.font.size = Pt(10)
        run.font.name = TNR
    return p


def add_figure_placeholder(doc, caption: str):
    """Bordered box + italic gray placeholder text, then a plain caption below."""
    box_p = doc.add_paragraph()
    box_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    box_p.paragraph_format.space_before = Pt(6)
    box_p.paragraph_format.space_after = Pt(2)
    add_box_border(box_p)
    run = box_p.add_run("[Insert Figure Here — Center within column]")
    run.italic = True
    run.font.size = Pt(9)
    run.font.name = TNR
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_before = Pt(2)
    cap_p.paragraph_format.space_after = Pt(6)
    cap_run = cap_p.add_run(caption)
    cap_run.font.size = Pt(9)
    cap_run.font.name = TNR


# ---------- main builder ----------

def build_ieee_docx(draft: PaperDraft, output_path: str):
    doc = Document()
    set_default_font(doc)
    set_margins(doc.sections[0])

    # --- Title: 24pt, NOT bold, centered ---
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_after = Pt(6)
    title_run = title_para.add_run(draft.title)
    title_run.font.size = Pt(24)
    title_run.font.name = TNR
    # title is intentionally NOT bold, matching the template

    # --- Abstract: bold+italic label, plain (non-italic) body ---
    abstract_para = doc.add_paragraph()
    abstract_para.paragraph_format.space_before = Pt(6)
    label_run = abstract_para.add_run("Abstract—")
    label_run.bold = True
    label_run.italic = True
    label_run.font.size = Pt(10)
    label_run.font.name = TNR
    body_run = abstract_para.add_run(draft.abstract)
    body_run.font.size = Pt(10)
    body_run.font.name = TNR

    # --- Keywords: italic-only label, plain body ---
    keywords_para = doc.add_paragraph()
    keywords_para.paragraph_format.space_before = Pt(4)
    keywords_para.paragraph_format.space_after = Pt(6)
    kw_label = keywords_para.add_run("Keywords — ")
    kw_label.italic = True
    kw_label.font.size = Pt(10)
    kw_label.font.name = TNR
    kw_body = keywords_para.add_run(", ".join(draft.keywords))
    kw_body.font.size = Pt(10)
    kw_body.font.name = TNR

    # --- Divider rule before the two-column body ---
    divider_para = doc.add_paragraph()
    divider_para.paragraph_format.space_after = Pt(6)
    add_bottom_border(divider_para)

    # --- Switch to two-column layout (0.25in gutter, matches template) ---
    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_margins(body_section)
    set_two_columns(body_section, num_cols=2, spacing=360)

    # --- Sections, Roman numeral headings ---
    roman_numerals = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X",
                       "XI", "XII", "XIII", "XIV", "XV"]
    for idx, sec in enumerate(draft.sections):
        heading_text = f"{roman_numerals[idx]}. {sec.section_title}"
        add_section_heading(doc, heading_text)
        add_body_paragraph(doc, sec.body)

        for fig in sec.figures_or_tables_needed:
            add_figure_placeholder(doc, f"Fig. {idx + 1}. {fig}")

    # --- References (unnumbered heading, matches template) ---
    add_section_heading_unnumbered = lambda text: add_section_heading(doc, text)
    add_section_heading_unnumbered("REFERENCES")

    for ref in draft.references:
        ref_para = doc.add_paragraph()
        ref_para.paragraph_format.left_indent = Inches(0.25)
        ref_para.paragraph_format.first_line_indent = Inches(-0.25)  # hanging indent
        ref_para.paragraph_format.space_after = Pt(2)  # 40 twips

        prefix_run = ref_para.add_run(f"[{ref.ieee_citation_number}] {ref.authors}, ")
        prefix_run.font.size = Pt(9)
        prefix_run.font.name = TNR

        title_run = ref_para.add_run(f'"{ref.title}," ')
        title_run.font.size = Pt(9)
        title_run.font.name = TNR

        venue_run = ref_para.add_run(f"{ref.venue}")
        venue_run.italic = True          # real italics, not literal \textit{}
        venue_run.font.size = Pt(9)
        venue_run.font.name = TNR

        year_run = ref_para.add_run(f", {ref.year}.")
        year_run.font.size = Pt(9)
        year_run.font.name = TNR

    doc.save(output_path)
    print(f"✓ Saved IEEE-format document to {output_path}")